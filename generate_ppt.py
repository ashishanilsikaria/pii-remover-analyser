import math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR


def estimate_row_content_lines(
    record, chars_per_line_desc=55, chars_per_line_findings=60
):

    heading_lines = math.ceil(len(record[2]) / chars_per_line_desc)
    desc_lines = math.ceil(len(record[3]) / chars_per_line_desc)

    findings_lines = 0
    for finding in record[4]:
        findings_lines += math.ceil(len(finding) / chars_per_line_findings)

    return max(heading_lines + desc_lines, findings_lines) + 2


def create_presentation(data, output_filename):

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    slide_layout_title = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout_title)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Automating File Cleansing and Analysis Leveraging AI"
    subtitle.text = (
        "Case Study Solution Presentation" "\nGroup Number: __" "\nTeam Members: __"
    )

    slide_layout_content = prs.slide_layouts[1]
    static_slides_info = [
        {
            "title": "Case study Overview",
            "content": "Populate this slide with your interpretation of the case study.",
        },
        {
            "title": "Approach Taken",
            "content": "Populate this slide with the strategy you used to address the case challenges.",
        },
        {
            "title": "Functional Design Diagram",
            "content": "Populate this slide with the functional design diagram of your application...",
        },
    ]

    for info in static_slides_info:
        slide = prs.slides.add_slide(slide_layout_content)
        slide.shapes.title.text = info["title"]
        slide.placeholders[1].text = info["content"]

    ppt_headers = ["File Name", "File Type", "File Description", "Key Findings"]

    MAX_CONTENT_LINES_PER_SLIDE = 18

    rows_to_process = data[:]

    empty_slide = prs.slide_layouts[5]
    first_table_slide = True

    while rows_to_process:
        slide = prs.slides.add_slide(empty_slide)
        if first_table_slide:
            slide.shapes.title.text = "File Analysis Sample Output"
            first_table_slide = False
        else:
            slide.shapes.title.text = ""

        current_slide_rows = []
        current_slide_lines = 0

        while rows_to_process:
            next_row = rows_to_process[0]
            estimated_lines = estimate_row_content_lines(next_row)

            if (
                not current_slide_rows
                or (current_slide_lines + estimated_lines)
                <= MAX_CONTENT_LINES_PER_SLIDE
            ):
                current_slide_lines += estimated_lines
                current_slide_rows.append(rows_to_process.pop(0))
            else:
                break

        num_rows_in_table = len(current_slide_rows) + 1

        side_margin_emu = int(Inches(0.2))
        top_padding_emu = int(Inches(1.0))
        bottom_margin_emu = int(Inches(0.25))

        slide_w_emu = int(prs.slide_width)
        slide_h_emu = int(prs.slide_height)
        title_shape = slide.shapes.title
        title_bottom_emu = int(title_shape.top + title_shape.height)

        table_left = Emu(side_margin_emu)
        table_top = Emu(title_bottom_emu + top_padding_emu)
        table_width = Emu(slide_w_emu - 2 * side_margin_emu)
        table_height = Emu(
            slide_h_emu - (title_bottom_emu + top_padding_emu) - bottom_margin_emu
        )

        table = slide.shapes.add_table(
            num_rows_in_table, 4, table_left, table_top, table_width, table_height
        ).table

        table.columns[0].width = Inches(1.4)
        table.columns[1].width = Inches(0.8)
        table.columns[2].width = Inches(3.6)
        table.columns[3].width = Inches(3.7)

        for col_idx, header_text in enumerate(ppt_headers):
            cell = table.cell(0, col_idx)
            cell.text = header_text
            p = cell.text_frame.paragraphs[0]
            p.font.bold = True
            p.font.size = Pt(12)
            p.alignment = PP_ALIGN.CENTER
            cell.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE

        header_height = Inches(0.4)
        table.rows[0].height = header_height

        data_rows = current_slide_rows
        row_line_counts = [estimate_row_content_lines(r) for r in data_rows]
        available_height_emu = int(table_height) - int(header_height)
        min_row_height_emu = int(Inches(0.3))
        per_line_emu = int(Pt(14))
        desired_heights = [
            max(min_row_height_emu, int(per_line_emu * lc)) for lc in row_line_counts
        ]
        total_desired = sum(desired_heights)
        if total_desired > available_height_emu and total_desired > 0:
            scale = available_height_emu / total_desired
            actual_heights = [
                max(min_row_height_emu, int(h * scale)) for h in desired_heights
            ]
        else:
            actual_heights = desired_heights

        for idx, h_emu in enumerate(actual_heights, start=1):
            table.rows[idx].height = Emu(h_emu)

        table_total_height = int(header_height) + sum(int(h) for h in actual_heights)
        table_shape = slide.shapes[-1]  # last added shape is the table
        try:
            table_shape.height = Emu(table_total_height)
        except Exception:
            pass

        for row_idx, record in enumerate(current_slide_rows, start=1):

            table.cell(row_idx, 0).text = record[0]
            table.cell(row_idx, 1).text = record[1]
            for col_idx in [0, 1]:
                cell = table.cell(row_idx, col_idx)
                cell.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
                cell.text_frame.paragraphs[0].font.size = Pt(12)

            cell = table.cell(row_idx, 2)
            cell.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
            tf = cell.text_frame
            tf.clear()
            tf.margin_top = Inches(0.02)
            tf.margin_bottom = Inches(0.02)
            tf.margin_left = Inches(0.03)
            tf.margin_right = Inches(0.03)

            p_h = tf.paragraphs[0]
            p_h.text = record[2]
            p_h.font.bold = True
            p_h.font.size = Pt(12)
            p_h.space_before = Pt(0)
            p_h.space_after = Pt(0)

            p_d = tf.add_paragraph()
            p_d.text = record[3]
            p_d.font.size = Pt(12)
            p_d.space_before = Pt(0)
            p_d.space_after = Pt(0)

            cell = table.cell(row_idx, 3)
            cell.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
            tf = cell.text_frame
            tf.clear()
            tf.margin_top = Inches(0.02)
            tf.margin_bottom = Inches(0.02)
            tf.margin_left = Inches(0.03)
            tf.margin_right = Inches(0.03)

            if record[4]:
                p_f = tf.paragraphs[0]
                p_f.text = f"• {record[4][0]}"
                p_f.font.size = Pt(12)
                p_f.space_before = Pt(0)
                p_f.space_after = Pt(0)

                for finding in record[4][1:]:
                    p_fb = tf.add_paragraph()
                    p_fb.text = f"• {finding}"
                    p_fb.font.size = Pt(12)
                    p_fb.space_before = Pt(0)
                    p_fb.space_after = Pt(0)
            else:

                tf.paragraphs[0].text = ""

    prs.save(output_filename)


if __name__ == "__main__":
    sample_data = [
        [
            "File_001.png",
            ".png",
            "Access Control System (Card Reader)",
            "The image depicts an individual using an RFID/NFC-based access control card to gain entry through a card reader, likely to a restricted area such as an IDF/Electrical room. The system is designed for physical security and access management.",
            [
                "The use of a physical access card with a photo ID enhances security by allowing visual verification of the cardholder.",
                "Access control systems like this provide granular control over who can enter specific areas and when, improving overall physical security.",
                "The room designation 'IDF/ELECTRICAL' indicates a critical infrastructure space, underscoring the importance of restricted access.",
                "Potential vulnerabilities include card cloning, loss or theft of access cards, or tailgating if not properly enforced.",
            ],
        ],
        [
            "File_002.png",
            ".png",
            "Biometric Access Control System",
            "The image displays a multi-factor authentication access control system featuring a keypad, a fingerprint scanner, and a camera, likely used for secure entry into a building or restricted area.",
            [
                "The presence of a fingerprint scanner indicates biometric authentication, which offers a higher level of security compared to traditional key or card systems.",
                "The keypad suggests an additional layer of security, possibly requiring a PIN in combination with biometrics (multi-factor authentication).",
                "The integrated camera could be used for facial recognition or recording individuals attempting to gain access, adding an extra layer of surveillance and identity verification.",
                "The system enhances physical security by controlling and monitoring entry points, deterring unauthorized access and providing an audit trail of entries.",
            ],
        ],
        [
            "File_003.png",
            ".png",
            "Visitors Log Book",
            "The image shows a traditional paper-based visitors log book, used to record visitor details, reasons for visit, and entry/exit times. It highlights manual data entry and the physical storage of visitor information.",
            [
                "Privacy concerns exist as visitor names and reasons for visit are visible to subsequent visitors.",
                "Data accuracy can be compromised due to illegible handwriting or human error during manual entry.",
                "Lack of advanced security features like access control, audit trails, or real-time alerts.",
                "Physical log books are susceptible to damage, loss, or unauthorized alteration, and do not offer robust data backup.",
            ],
        ],
    ]
    create_presentation(sample_data, "case_study_presentation.pptx")
